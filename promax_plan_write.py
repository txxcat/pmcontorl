# -*- coding: utf-8 -*-
"""把测试结果写成 **PROMAX 仪器导出的 XML**（下游软件只认这个格式）。

为什么需要
----------
仪器自带的 PC 软件 / 第三方报表工具只吃它自己导出的 `.XML`。
本程序的结果是 xlsx / csv，要接这类程序就得转一道。

格式（按真机导出的 version="1.3" 结构）::

    <?xml version="1.0" encoding="utf-8"?>
    <POINT name="tw260928" date="2026-09-28" time="15:12:49" >
      <GPS latitude="0.0000000" longitude="0.0000000" altitude="0.00" locked="false" />
      <CHANNEL-SET band="TERRESTRIAL" date=".." equipment="RANGERNeo 4"
                   name="表名" time=".." version="1.3" >
        <CHANNEL name="171-DS06" frequency="171.00" units="MHz" date=".." time=".." >
          <MEASUREMENTS>
            <DVB-C>
              <PARAMETERS><SYMBOLRATE value="6875" units="kSymb/s" /></PARAMETERS>
              <MEASURES>
                <POWER value="82.5" range="IN" units="dBuV" />
                <CN value="38.9" range="IN" units="dB" />
                <MER value="42.2" range="IN" units="dB" />
                <OFFSET value="-2.9" range="IN" units="kHz" />
              </MEASURES>
              <STATUS value="MPEG2 TS locked" locked="LOCKED" />
            </DVB-C>
          </MEASUREMENTS>
        </CHANNEL>
      </CHANNEL-SET>
    </POINT>

要点（都按样本实测核对过）：
  · 文件编码 utf-8 **无 BOM**、行尾 **CRLF**、缩进 **Tab**、属性用双引号
  · 射频频道用 ``frequency``（调谐频率，模拟频道就是图像载频）；
    卫星频道**没有 frequency**，用 ``downlink``（下行） + ``sat_band="LOW"``
    —— 我们的结果里卫星存的是**中频**，所以要靠频道表回推下行
  · 制式块：模拟 → ``<UNKNOWN>``（真机也不给模拟频道定标准），
    数字 → ``<DVB-C>`` / ``<DVB-S>`` / ``<DVB-S2>``
  · ``STATUS``：锁定 → ``value="MPEG2 TS locked" locked="LOCKED"``，
    否则 → ``value="No signal received" locked="UNLOCKED"``
  · 真机还会写 CONSTELLATION / CODERATE / CBER / LM 等，**我们测不到就不写**
    （宁可缺字段，也不能编数据）

用法
----
  python promax_plan_write.py 结果.xlsx            # 生成同名 .XML
  python promax_plan_write.py 结果.csv -o out.XML
"""
import csv
import datetime
import os
import re
import sys

EQUIPMENT = "RANGERNeo 4"
VERSION = "1.3"

# 卫星下行 ↔ 中频的本振（与 promax_plan_dump 的口径一致）
LO_C = 5150.0       # C 波段：中频 = 本振 − 下行
LO_KU = 11300.0     # Ku 波段：中频 = 下行 − 本振

# 结果文件表头 → 内部分键（表头里的单位括号会先被去掉）
HEADER_KEYS = {
    "序号": "idx", "频段": "band", "频道表": "plan", "频道": "ch",
    "频率": "freq_mhz", "制式": "signal", "LOCK": "LOCKED",
    "LEVEL": "LEVEL", "CN": "CN", "POWER": "POWER", "MER": "MER",
    "OFFSET": "OFFSET_khz",       # 结果文件里已经是 kHz
    "SR": "symb_rate", "尝试": "attempts",
}


# --------------------------------------------------------------- 基础工具
def _num(v):
    """'82.5' → 82.5；空 / 非数字 → None"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _f(v, nd):
    """数字格式化；None → None（调用方据此跳过该元素）"""
    x = _num(v)
    return None if x is None else ("%%.%df" % nd) % x


def _sr_text(v):
    """符号率：整数就不带小数点（真机写 6875，不是 6875.0）"""
    x = _num(v)
    if x is None:
        return None
    return "%d" % int(round(x)) if abs(x - round(x)) < 1e-6 else "%.1f" % x


def _attr(v):
    """XML 属性值转义（含双引号 / & < >），返回带引号的串"""
    s = "" if v is None else str(v)
    s = (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
          .replace('"', "&quot;"))
    return '"%s"' % s


def standard_tag(signal):
    """制式 → 元素名。模拟（连续载波）真机写 UNKNOWN。"""
    s = str(signal or "").strip().upper()
    if not s or s.startswith("ANALOG"):
        return "UNKNOWN"
    for tag in ("DVB-S2", "DVB-S", "DVB-C", "DVB-T2", "DVB-T"):
        if s.startswith(tag):
            return tag
    return "UNKNOWN"


def default_promax_path(result_path):
    """结果文件 → 同名 .XML（与结果同目录）"""
    return os.path.splitext(result_path)[0] + ".XML"


# --------------------------------------------------------------- 组装频道
def table_downlinks(plans_dir, plan_name):
    """从频道表目录里找同名表 → {频道名: 下行MHz}（找不到返回 {}）"""
    if not plans_dir:
        return {}
    try:
        from promax_plan_dump import load_plans
        ter, sat = load_plans(plans_dir)
    except Exception:
        return {}
    want = str(plan_name).strip()
    for p in list(sat) + list(ter):
        names = (str(p.get("name", "")).strip(),
                 str(p.get("exact_file", "")).strip(),
                 str(p.get("file", "")).strip())
        if want in names:
            out = {}
            for c in p.get("channels") or []:
                if c.get("freq_mhz") is not None:
                    out[str(c.get("name", "")).strip()] = c["freq_mhz"]
            return out
    return {}


def build_sets(results, downlink_map=None):
    """{表名: [row,...]} → [(表名, 'TER'|'SAT', [频道 dict, ...]), ...]

    `downlink_map`：{(表名, 频道名): 下行MHz}。卫星结果里存的是**中频**，
    转成 Promax 必须回推下行；传了就用频道表里的真值，没传就按本振反推。
    """
    order, groups = [], {}
    for plan, rows in (results or {}).items():
        for r in rows or []:
            band = str(r.get("band") or "TER").upper()
            band = "SAT" if band.startswith("SAT") else "TER"
            key = (plan, band)
            if key not in groups:
                groups[key] = []
                order.append(key)
            groups[key].append(_channel(r, plan, band, downlink_map))
    return [(plan, band, groups[(plan, band)]) for plan, band in order]


def _channel(r, plan, band, downlink_map):
    """一条结果 → 频道 dict"""
    freq = _num(r.get("freq_mhz"))
    off_khz = r.get("OFFSET_khz")
    if off_khz is None:                        # 内存里的 row：OFFSET 是 Hz
        ov = _num(r.get("OFFSET"))
        off_khz = None if ov is None else ov / 1000.0
    power = _num(r.get("POWER"))
    if power is None:
        power = _num(r.get("LEVEL"))           # 模拟制式的电平在 LEVEL 上
    ch = {"name": str(r.get("ch") or "").strip(),
          "band": band, "freq": freq,
          "std": standard_tag(r.get("signal")),
          "locked": str(r.get("LOCKED") or "").strip().upper() == "ON",
          "power": power, "cn": _num(r.get("CN")), "mer": _num(r.get("MER")),
          "offset_khz": off_khz, "sr": r.get("symb_rate")}
    if band == "SAT":
        dl = None
        if downlink_map and freq is not None:
            dl = downlink_map.get((str(plan).strip(), ch["name"]))
        if dl is None and freq is not None:
            # 频道名本身就是下行（现场命名习惯："3780" / "12395"）→ 直接用
            nm = _num(ch["name"])
            if nm is not None and (3400 <= nm <= 4300 or 10500 <= nm <= 13250):
                dl = nm
        if dl is None and freq is not None:
            # 最后兜底：C 波段下行 = 本振 − 中频；Ku = 下行 − 本振。
            # ⚠️ 只看中频**分不出** C/Ku（两者中频都落在 950~2150），
            #    所以优先靠频道表（downlink_map）或频道名，这里只是保底。
            dl = (LO_C - freq) if freq <= 2150 else (LO_KU + freq)
        ch["downlink"] = dl
    return ch


# --------------------------------------------------------------- 写 XML
def _channel_lines(ch, date_s, time_s):
    """一个 <CHANNEL> 块（含 MEASUREMENTS）"""
    L = []
    if ch["band"] == "SAT":
        head = ('\t\t<CHANNEL name=%s downlink=%s sat_band="LOW" units="MHz" '
                'date="%s" time="%s" >'
                % (_attr(ch["name"]), _attr(_f(ch.get("downlink"), 2)),
                   date_s, time_s))
    else:
        head = ('\t\t<CHANNEL name=%s frequency=%s units="MHz" '
                'date="%s" time="%s" >'
                % (_attr(ch["name"]), _attr(_f(ch.get("freq"), 2)),
                   date_s, time_s))
    L.append(head)
    L.append("\t\t\t<MEASUREMENTS>")
    std = ch["std"]
    L.append("\t\t\t\t<%s>" % std)

    sr = _sr_text(ch.get("sr"))
    if sr is not None and std != "UNKNOWN":
        L.append("\t\t\t\t\t<PARAMETERS>")
        L.append('\t\t\t\t\t\t<SYMBOLRATE value="%s" units="kSymb/s" />' % sr)
        L.append("\t\t\t\t\t</PARAMETERS>")

    L.append("\t\t\t\t\t<MEASURES>")
    # 模拟制式的 POWER 真机标 range="UNKNOWN"；数字是 "IN"
    rng = "UNKNOWN" if std == "UNKNOWN" else "IN"
    p = _f(ch.get("power"), 1)
    if p is not None:
        L.append('\t\t\t\t\t\t<POWER value="%s" range="%s" units="dBuV" />'
                 % (p, rng))
    cn = _f(ch.get("cn"), 1)
    if cn is not None:
        L.append('\t\t\t\t\t\t<CN value="%s" range="IN" units="dB" />' % cn)
    mer = _f(ch.get("mer"), 1)
    if mer is not None:
        L.append('\t\t\t\t\t\t<MER value="%s" range="IN" units="dB" />' % mer)
    off = _f(ch.get("offset_khz"), 1)
    if off is not None:
        L.append('\t\t\t\t\t\t<OFFSET value="%s" range="IN" units="kHz" />'
                 % off)
    L.append("\t\t\t\t\t</MEASURES>")

    if ch["locked"]:
        L.append('\t\t\t\t\t<STATUS value="MPEG2 TS locked" '
                 'locked="LOCKED" />')
    else:
        L.append('\t\t\t\t\t<STATUS value="No signal received" '
                 'locked="UNLOCKED" />')
    L.append("\t\t\t\t</%s>" % std)
    L.append("\t\t\t</MEASUREMENTS>")
    L.append("\t\t</CHANNEL>")
    return L


def write_promax(path, sets, name=None, equipment=EQUIPMENT, when=None):
    """写 Promax XML。`sets` = build_sets() 的返回值。

    返回写出的路径。空数据会抛 ValueError（别生成一个没有频道的空文件）。
    """
    total = sum(len(chs) for _n, _b, chs in sets or [])
    if not total:
        raise ValueError("没有可写入的频道记录")
    d = when or datetime.datetime.now()
    date_s, time_s = d.strftime("%Y-%m-%d"), d.strftime("%H:%M:%S")
    if not name:
        name = os.path.splitext(os.path.basename(path))[0]

    L = ['<?xml version="1.0" encoding="utf-8"?>',
         '<POINT name=%s date="%s" time="%s" >' % (_attr(name), date_s, time_s),
         '\t<GPS latitude="0.0000000" longitude="0.0000000" '
         'altitude="0.00" locked="false" />']
    for set_name, band, chans in sets:
        L.append('\t<CHANNEL-SET band="%s" date="%s" equipment=%s name=%s '
                 'time="%s" version="%s" >'
                 % ("SATELLITE" if band == "SAT" else "TERRESTRIAL",
                    date_s, _attr(equipment), _attr(set_name), time_s, VERSION))
        for ch in chans:
            L.extend(_channel_lines(ch, date_s, time_s))
        L.append("\t</CHANNEL-SET>")
    L.append("</POINT>")

    d0 = os.path.dirname(os.path.abspath(path))
    if d0 and not os.path.isdir(d0):
        os.makedirs(d0, exist_ok=True)
    # newline="" → 自己控制 CRLF（默认文本模式会把 \r\n 再翻成 \r\r\n）
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write("\r\n".join(L) + "\r\n")
    return path


# --------------------------------------------------------------- 读结果文件
def _norm_header(h):
    """'OFFSET(kHz)' → 'OFFSET'；'频率(MHz)' → '频率'"""
    return re.sub(r"[（(].*?[)）]", "", str(h or "")).strip()


def read_result_file(path):
    """读本程序导出的结果文件（xlsx / csv）→ {表名: [row,...]}

    row 的键与内存里的结果行一致，只是 OFFSET 放在 `OFFSET_khz`（文件里已是 kHz）。
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        headers, table = _read_xlsx(path)
    elif ext == ".csv":
        headers, table = _read_csv(path)
    else:
        raise ValueError("只支持 xlsx / csv：%s" % path)

    keys = []
    for h in headers:
        k = HEADER_KEYS.get(_norm_header(h))
        keys.append(k)
    if "ch" not in keys or "freq_mhz" not in keys:
        raise ValueError("表头不像本程序的结果文件（缺少「频道」或「频率」列）：\n%s"
                         % " | ".join(str(h) for h in headers))

    out = {}
    for cells in table:
        r = {}
        for k, v in zip(keys, cells):
            if k:
                r[k] = v
        if not str(r.get("ch") or "").strip():
            continue
        band = str(r.get("band") or "").upper()
        r["band"] = "SAT" if band.startswith("SAT") else "TER"
        plan = str(r.get("plan") or "").strip() or "测试结果"
        r["plan"] = plan
        out.setdefault(plan, []).append(r)
    if not out:
        raise ValueError("文件里没有数据行：%s" % path)
    return out


def _read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        try:
            wb.close()
        except Exception:
            pass
    if not rows:
        raise ValueError("空文件：%s" % path)
    return [c for c in rows[0]], rows[1:]


def _read_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        rd = list(csv.reader(f))
    if not rd:
        raise ValueError("空文件：%s" % path)
    return rd[0], rd[1:]


# --------------------------------------------------------------- 一步到位
def convert_result_file(src, dst=None, downlink_map=None, plans_dir=None):
    """结果文件 → Promax XML。返回 (输出路径, 频道数)。

    `plans_dir`：频道表目录。结果文件里卫星只有中频，靠同名频道表才能拿到
    真实下行；没有表时退回"频道名即下行"或本振推算。
    """
    results = read_result_file(src)
    dmap = dict(downlink_map or {})
    if plans_dir:
        for plan, rows in results.items():
            if not any(str(r.get("band", "")).upper().startswith("SAT")
                       for r in rows):
                continue
            for k, v in table_downlinks(plans_dir, plan).items():
                dmap.setdefault((str(plan).strip(), k), v)
    sets = build_sets(results, dmap)
    dst = dst or default_promax_path(src)
    write_promax(dst, sets, name=os.path.splitext(os.path.basename(dst))[0])
    return dst, sum(len(chs) for _n, _b, chs in sets)


def write_promax_for_results(results, path, downlink_map=None, name=None):
    """内存里的结果（{表名: [row,...]}）→ Promax XML。返回 (路径, 频道数)。"""
    sets = build_sets(results, downlink_map)
    write_promax(path, sets, name=name)
    return path, sum(len(chs) for _n, _b, chs in sets)


# --------------------------------------------------------------- CLI
def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv or argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0
    src = argv[0]
    dst = None
    if "-o" in argv:
        i = argv.index("-o")
        if i + 1 < len(argv):
            dst = argv[i + 1]
    try:
        out, n = convert_result_file(src, dst)
    except Exception as e:
        print("转换失败: %s" % e)
        return 2
    print("已生成: %s（%d 个频道）" % (out, n))
    return 0


if __name__ == "__main__":
    sys.exit(main())
